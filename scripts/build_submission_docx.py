from __future__ import annotations

import re
import unicodedata
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
OUTPUT = PAPER / "submission"


def set_font(run, name: str = "Times New Roman", size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def plain_latex(text: str) -> str:
    replacements = {
        r"\mathbb{R}": "ℝ",
        r"\times": "×",
        r"\in": "∈",
        r"\phi": "φ",
        r"\nu": "ν",
        r"\sigma": "σ",
        r"\tau": "τ",
        r"\ldots": "…",
        r"\mathrm{ITR}": "ITR",
        r"\log_2": "log₂",
        r"\sqrt": "√",
        r"\mathsf T": "T",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = text.replace("\\left", "").replace("\\right", "")
    text = text.replace("\\,", " ").replace("\\!", "")
    text = text.replace("{", "").replace("}", "")
    return text


def decode_bibtex_text(text: str) -> str:
    """Decode common BibTeX accent commands without a heavyweight parser."""
    text = text.replace(r"\i", "i").replace(r"\j", "j")
    combining = {
        "'": "\u0301",
        "`": "\u0300",
        "^": "\u0302",
        '"': "\u0308",
        "~": "\u0303",
        "=": "\u0304",
        "c": "\u0327",
    }

    def replace_accent(match: re.Match[str]) -> str:
        return unicodedata.normalize("NFC", match.group(2) + combining[match.group(1)])

    text = re.sub(r"\\(['`^\"~=c])\{?([A-Za-z])\}?", replace_accent, text)
    text = text.replace(r"\&", "&").replace(r"\_", "_").replace(r"\%", "%")
    return text.replace("{", "").replace("}", "")


def parse_bibliography(path: Path) -> dict[str, dict[str, str]]:
    entries: dict[str, dict[str, str]] = {}
    current: dict[str, str] | None = None
    key = ""
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        match = re.match(r"@\w+\{([^,]+),", line)
        if match:
            key = match.group(1).strip()
            current = {}
            entries[key] = current
            continue
        if current is None:
            continue
        field = re.match(r"(\w+)\s*=\s*\{(.*)\}\s*,?\s*$", line)
        if field:
            value = decode_bibtex_text(field.group(2))
            current[field.group(1).lower()] = value
    return entries


def cite_order(text: str) -> tuple[str, list[str]]:
    order: list[str] = []

    def replace(match: re.Match[str]) -> str:
        keys = []
        for part in match.group(1).split(";"):
            candidate = part.strip()
            if candidate.startswith("@"):
                key = candidate[1:].strip()
                if key not in order:
                    order.append(key)
                keys.append(str(order.index(key) + 1))
        return "[" + ", ".join(keys) + "]"

    return re.sub(r"\[([^\]]*@[^\]]+)\]", replace, text), order


def add_inline(paragraph, text: str) -> None:
    token = re.compile(r"(\*\*.*?\*\*|`.*?`|\\\(.*?\\\))")
    position = 0
    for match in token.finditer(text):
        if match.start() > position:
            set_font(paragraph.add_run(text[position:match.start()]), size=10.5)
        value = match.group(0)
        if value.startswith("**"):
            run = paragraph.add_run(value[2:-2])
            run.bold = True
            set_font(run, size=10.5)
        elif value.startswith("`"):
            run = paragraph.add_run(value[1:-1])
            set_font(run, name="Consolas", size=9.0)
        else:
            set_font(paragraph.add_run(plain_latex(value[2:-2])), size=10.5)
        position = match.end()
    if position < len(text):
        set_font(paragraph.add_run(text[position:]), size=10.5)


def render_equation(latex: str, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    normalized = latex.replace(r"\mathsf T", "T")
    fig = plt.figure(figsize=(6.6, 0.8), dpi=300)
    fig.patch.set_alpha(0)
    fig.text(0.5, 0.5, f"${normalized}$", ha="center", va="center", fontsize=12)
    fig.savefig(path, dpi=300, transparent=True, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if row_index == 0:
                shade_cell(cell, "263849")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            add_inline(paragraph, value)
            for run in paragraph.runs:
                set_font(run, size=8.2)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
        if row_index > 0 and row_index % 2 == 0:
            for cell in table.rows[row_index].cells:
                shade_cell(cell, "F3F6F8")
    set_repeat_table_header(table.rows[0])
    document.add_paragraph()


def format_reference(index: int, entry: dict[str, str]) -> str:
    authors = entry.get("author", "Unknown author").replace(" and ", ", ")
    title = entry.get("title", "Untitled")
    year = entry.get("year", "")
    journal = entry.get("journal", entry.get("booktitle", entry.get("publisher", "")))
    volume = entry.get("volume", "")
    number = entry.get("number", "")
    pages = entry.get("pages", "").replace("--", "–")
    doi = entry.get("doi", "")
    source = journal
    if volume:
        source += f" {volume}"
    if number:
        source += f"({number})"
    if pages:
        source += f":{pages}"
    if doi:
        source += f". https://doi.org/{doi}"
    return f"{index}. {authors}. {title}. {source} ({year})."


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, before, after in (
        ("Title", 18, 0, 12),
        ("Heading 1", 14, 12, 5),
        ("Heading 2", 12, 9, 4),
        ("Heading 3", 10.5, 7, 3),
    ):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        # The host Word template may carry an accent-coloured bottom border on
        # its built-in Title style. Remove that inherited decoration so the
        # submission proof is deterministic across machines.
        if style_name == "Title":
            paragraph_properties = style._element.get_or_add_pPr()
            border = paragraph_properties.find(qn("w:pBdr"))
            if border is not None:
                paragraph_properties.remove(border)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("HarmonicFoldNet preprint proof")
    set_font(run, name="Arial", size=8)
    run.font.color.rgb = RGBColor(100, 100, 100)


def build_document(source: Path, destination: Path, bibliography: Path | None = None) -> None:
    raw = source.read_text(encoding="utf-8")
    if bibliography is not None:
        raw, citation_keys = cite_order(raw)
        references = parse_bibliography(bibliography)
    else:
        citation_keys = []
        references = {}

    document = Document()
    configure_document(document)
    lines = raw.splitlines()
    index = 0
    equation_index = 0
    skip_references = False
    while index < len(lines):
        stripped = lines[index].strip()
        if skip_references:
            index += 1
            continue
        if stripped == "## References" and bibliography is not None:
            skip_references = True
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, stripped[2:])
            for run in paragraph.runs:
                set_font(run, name="Arial", size=18)
            index += 1
            continue
        if stripped.startswith("### "):
            document.add_paragraph(stripped[4:], style="Heading 2")
            index += 1
            continue
        if stripped.startswith("## "):
            document.add_paragraph(stripped[3:], style="Heading 1")
            index += 1
            continue
        if stripped.startswith("!["):
            image_match = re.match(r"!\[[^\]]*\]\(([^)]+)\)", stripped)
            if image_match:
                image_path = source.parent / image_match.group(1)
                png_path = image_path.with_suffix(".png")
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                paragraph.add_run().add_picture(str(png_path), width=Inches(6.65))
            index += 1
            continue
        if stripped == r"\[":
            equation_lines: list[str] = []
            index += 1
            while index < len(lines) and lines[index].strip() != r"\]":
                equation_lines.append(lines[index].strip())
                index += 1
            equation_index += 1
            equation_path = OUTPUT / "qa_equations" / f"equation-{equation_index}.png"
            render_equation(" ".join(equation_lines), equation_path)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(equation_path), width=Inches(5.8))
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(document, table_lines)
            continue
        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith("#")
                or candidate.startswith("![")
                or candidate.startswith("|")
                or candidate.startswith("- ")
                or candidate == r"\["
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        text = " ".join(paragraph_lines)
        add_inline(paragraph, text)
        if text.startswith("**Figure ") or text.startswith("**Table "):
            paragraph.paragraph_format.keep_with_next = False
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(7)
        elif text.startswith(
            (
                "**Authors:**",
                "**ORCID:**",
                "**Affiliation:**",
                "**Affiliations:**",
                "**Corresponding author:**",
            )
        ):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if bibliography is not None:
        document.add_paragraph("References", style="Heading 1")
        for ref_index, key in enumerate(citation_keys, start=1):
            if key not in references:
                raise KeyError(f"missing bibliography entry: {key}")
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.space_after = Pt(3)
            set_font(paragraph.add_run(format_reference(ref_index, references[key])), size=9)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> None:
    build_document(
        PAPER / "MANUSCRIPT_DRAFT_v1.md",
        OUTPUT / "HarmonicFoldNet_manuscript_proof.docx",
        PAPER / "references.bib",
    )
    build_document(
        PAPER / "SUPPLEMENTARY_INFORMATION.md",
        OUTPUT / "HarmonicFoldNet_supplement_proof.docx",
    )
    print(
        {
            "manuscript": str(OUTPUT / "HarmonicFoldNet_manuscript_proof.docx"),
            "supplement": str(OUTPUT / "HarmonicFoldNet_supplement_proof.docx"),
        }
    )


if __name__ == "__main__":
    main()
